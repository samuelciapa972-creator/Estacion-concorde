"""API (FastAPI) contra PostgreSQL real: éxito, validación y permisos de cada endpoint, y el flujo completo
simulador -> pendientes -> cliente -> emitir -> worker -> FACTURADO. Solo ProveedorSimulado."""

from __future__ import annotations

import io
import logging
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any

import jwt
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import ENTORNO_DIAN, PIN, SECRETO, Api, ajustes, clave, empresa, persona
from estacion import usuarios
from estacion.api import crear_app
from estacion.api.__main__ import SinConsulta
from estacion.api.ajustes import Ajustes, AjustesInvalidos
from estacion.facturacion.proveedor import ProveedorSimulado
from estacion.servicios.outbox import procesar_pendientes
from estacion.surtidores import SimuladorSurtidor, ingerir

pytestmark = pytest.mark.bd


# ------------------------------------------------------------------------------------------ arranque y salud
def test_salud(api):
    assert api.get("/salud").json() == {"estado": "ok", "base_de_datos": "ok", "migracion": "0004"}


def test_salud_con_la_base_sin_migraciones(esquema_vacio):
    with TestClient(crear_app(ajustes(esquema_vacio.dsn), entorno_dian=ENTORNO_DIAN)) as c:
        r = c.get("/salud")
    assert r.status_code == 503 and "make migrate" in r.json()["detail"]


@pytest.mark.parametrize(
    ("entorno", "mensaje"),
    [
        ({"JWT_SECRETO": SECRETO}, "DATABASE_URL"),
        ({"DATABASE_URL": "postgresql://x"}, "JWT_SECRETO"),
        ({"DATABASE_URL": "postgresql://x", "JWT_SECRETO": "corto"}, "JWT_SECRETO"),
        ({"DATABASE_URL": "postgresql://x", "JWT_SECRETO": SECRETO, "JWT_MINUTOS": "mucho"}, "JWT_MINUTOS"),
    ],
)
def test_sin_ajustes_la_api_no_arranca(entorno, mensaje):
    with pytest.raises(AjustesInvalidos, match=mensaje):
        Ajustes.desde_entorno(entorno)


def test_openapi_documenta_los_endpoints(api):
    rutas = api.get("/openapi.json").json()["paths"]
    for ruta in ("/auth/login", "/despachos/pendientes", "/ventas-manuales", "/clientes/buscar", "/clientes",
                 "/facturas", "/facturas/{factura_id}", "/reportes/diario", "/reportes/mensual", "/reportes/turno",
                 "/reportes/exportar", "/anomalias", "/salud", "/publico/registro", "/catalogos"):  # fmt: skip
        assert ruta in rutas


# ------------------------------------------------------------------------------------------ sesión
def test_ingreso_y_usuario_actual(api):
    r = api.post("/auth/login", {"usuario": api.bombero.upper(), "pin": PIN})
    assert r.status_code == 200 and r.json()["expira_en"] == 3600 and r.json()["usuario"]["rol"] == "BOMBERO"
    yo = api.cliente.get("/auth/yo", headers={"Authorization": f"Bearer {r.json()['access_token']}"})
    assert yo.json()["usuario"] == api.bombero


def test_pin_incorrecto_y_usuario_inexistente_responden_igual(api):
    malo = api.post("/auth/login", {"usuario": api.bombero, "pin": "0000"})
    nadie = api.post("/auth/login", {"usuario": "nadie", "pin": "0000"})
    assert malo.status_code == nadie.status_code == 401
    assert malo.json() == nadie.json() == {"detail": "usuario o PIN incorrectos"}


def test_bloqueo_tras_varios_pin_incorrectos(armar):
    api = armar(max_intentos_pin=3)
    for _ in range(3):
        assert api.post("/auth/login", {"usuario": api.bombero, "pin": "0000"}).status_code == 401
    assert api.post("/auth/login", {"usuario": api.bombero, "pin": PIN}).status_code == 423  # aun con el correcto
    acciones = [a for (a,) in api.e.datos.conn.execute("SELECT accion FROM auditoria ORDER BY id").fetchall()]
    assert acciones[-3:] == ["INGRESO_FALLIDO", "INGRESO_FALLIDO", "USUARIO_BLOQUEADO"]

    api.e.datos.conn.execute("UPDATE usuario SET bloqueado_hasta = now() - interval '1 second'")
    assert api.post("/auth/login", {"usuario": api.bombero, "pin": PIN}).status_code == 200


def test_un_ingreso_correcto_reinicia_los_intentos(armar):
    api = armar(max_intentos_pin=3)
    for _ in range(2):
        api.post("/auth/login", {"usuario": api.bombero, "pin": "0000"})
    api.entrar(api.bombero)
    for _ in range(2):
        api.post("/auth/login", {"usuario": api.bombero, "pin": "0000"})
    assert api.post("/auth/login", {"usuario": api.bombero, "pin": PIN}).status_code == 200


def test_limite_de_ingresos_por_equipo(armar):
    api = armar(login_por_minuto=2)
    for _ in range(2):
        api.post("/auth/login", {"usuario": "nadie", "pin": "0000"})
    assert api.post("/auth/login", {"usuario": api.bombero, "pin": PIN}).status_code == 429


@pytest.mark.parametrize(
    "cabeceras",
    [
        {},
        {"Authorization": "Bearer basura"},
        {
            "Authorization": "Bearer "
            + jwt.encode({"sub": "1", "exp": datetime.now(UTC) - timedelta(minutes=1)}, SECRETO)
        },
        {
            "Authorization": "Bearer "
            + jwt.encode({"sub": "1", "exp": datetime.now(UTC) + timedelta(hours=1)}, "x" * 40)
        },
    ],
    ids=["sin-token", "token-basura", "token-vencido", "firma-ajena"],
)
def test_sin_sesion_valida_no_hay_acceso(api, cabeceras):
    r = api.cliente.get("/despachos/pendientes", headers=cabeceras)
    assert r.status_code == 401 and r.headers["www-authenticate"] == "Bearer"


def test_desactivar_al_usuario_corta_su_sesion(api):
    cabeceras = api.entrar(api.bombero)
    api.e.datos.conn.execute("UPDATE usuario SET activo = FALSE WHERE usuario = %s", (api.bombero,))
    assert api.cliente.get("/auth/yo", headers=cabeceras).status_code == 401


# ------------------------------------------------------------------------------------------ pista
def test_catalogos(api):
    api.e.datos.despacho()
    r = api.get("/catalogos", api.bombero).json()
    assert [p["codigo"] for p in r["productos"]] == ["DIESEL"]
    assert r["medios_pago"] == ["EFECTIVO", "TARJETA"] and "NIT" in r["tipos_documento"]


def test_pendientes_con_dinero_y_galones_como_texto(api):
    despacho = api.e.datos.despacho()
    (p,) = api.get("/despachos/pendientes", api.bombero).json()
    assert (p["origen_tipo"], p["origen_id"]) == ("despacho", despacho)
    assert (p["volumen"], p["valor"], p["ppu"]) == ("1.861", "29757", "15990.00")
    assert api.get("/despachos/pendientes?lado=B", api.bombero).json() == []
    assert api.get("/despachos/pendientes?limite=0", api.bombero).status_code == 422


def test_venta_manual_por_api(api):
    api.e.datos.despacho()
    surtidor = api.e.datos.uno("SELECT id FROM surtidor")
    venta = {"surtidor_id": surtidor, "lado": "B", "producto": "diesel", "volumen": "3.120", "valor": "49889"}
    r = api.post("/ventas-manuales", venta, api.bombero)
    assert r.status_code == 201
    pendientes = api.get("/despachos/pendientes?lado=B", api.bombero).json()
    assert [(p["origen_tipo"], p["origen_id"], p["ppu"]) for p in pendientes] == [
        ("venta_manual", r.json()["id"], "15990.06")
    ]
    usuario = api.e.datos.uno("SELECT usuario_id FROM venta_manual WHERE id = %s", (r.json()["id"],))
    assert usuario == api.e.datos.uno("SELECT id FROM usuario WHERE usuario = %s", (api.bombero,))


@pytest.mark.parametrize(
    ("cambios", "en"),
    [
        ({"volumen": "3.1205"}, "volumen"),  # Pydantic: más de 3 decimales
        ({"valor": "-5"}, "valor"),
        ({"ppu": "16500"}, "no coincide"),  # el servicio: galones × precio no cuadra con el valor
        ({"producto": "KEROSENE"}, "producto desconocido"),
        ({"extra": 1}, "extra"),
    ],
)
def test_venta_manual_invalida(api, cambios, en):
    api.e.datos.despacho()
    surtidor = api.e.datos.uno("SELECT id FROM surtidor")
    venta = {"surtidor_id": surtidor, "lado": "B", "producto": "DIESEL", "volumen": "3.120", "valor": "49889"}
    r = api.post("/ventas-manuales", {**venta, **cambios}, api.bombero)
    assert r.status_code == 422 and en in r.text


# ------------------------------------------------------------------------------------------ clientes
def test_registrar_y_buscar_persona(api):
    r = api.post("/clientes", persona(), api.bombero)
    assert r.status_code == 201
    c = r.json()
    assert (c["nombres"], c["apellidos"], c["nombre_mostrar"]) == ("LUIS CARLOS", "ROJAS", "LUIS CARLOS ROJAS")
    assert c["email_enmascarado"] == "lu********@example.com"
    assert api.get("/clientes/buscar?tipo=CC&numero=1049612345", api.bombero).json() == c


def test_nit_muestra_la_razon_social_y_calcula_el_dv(api):
    c = api.post("/clientes", empresa(), api.bombero).json()
    assert (c["numero_documento"], c["digito_verificacion"]) == ("900123456", "8")
    assert c["nombre_mostrar"] == "TRANSPORTES DE PRUEBA S.A.S."  # la falla de Nexus: aquí siempre sale
    assert api.get("/clientes/buscar?tipo=NIT&numero=900123456", api.bombero).json()["nombre_mostrar"] == (
        "TRANSPORTES DE PRUEBA S.A.S."
    )


@pytest.mark.parametrize(
    ("entrada", "problema"),
    [
        (empresa(digito_verificacion="3"), "dígito de verificación incorrecto"),
        (persona(autoriza_tratamiento=False), "autorización de tratamiento"),
        (persona(email="sin-arroba"), "correo"),
        (persona(numero_documento="12AB"), "cédula"),
    ],
)
def test_cliente_invalido(api, entrada, problema):
    r = api.post("/clientes", entrada, api.bombero)
    assert r.status_code == 422 and problema in r.text
    assert api.e.datos.uno("SELECT count(*) FROM cliente WHERE numero_documento <> %s", ("x",)) == 1  # el del escenario


def test_cliente_repetido(api):
    primero = api.post("/clientes", persona(), api.bombero).json()
    r = api.post("/clientes", persona(nombres="OTRO"), api.bombero)
    assert r.status_code == 409 and r.json()["detail"]["cliente_id"] == primero["id"]


def test_buscar_cliente_inexistente_y_sin_sesion(api):
    assert api.get("/clientes/buscar?tipo=CC&numero=12345678", api.bombero).status_code == 404
    assert api.get("/clientes/buscar?tipo=CC&numero=12345678").status_code == 401
    assert api.get("/clientes/buscar?tipo=XX&numero=12345678", api.bombero).status_code == 422


def test_la_auditoria_no_guarda_datos_personales(api):
    api.post("/clientes", persona(), api.bombero)
    detalles = str(api.e.datos.conn.execute("SELECT detalle FROM auditoria").fetchall())
    for dato in ("1049612345", "ROJAS", "example.com"):
        assert dato not in detalles


# ------------------------------------------------------------------------------------------ registro público
def test_registro_publico_sin_sesion(api):
    r = api.post("/publico/registro", persona())
    assert r.status_code == 202
    assert api.e.datos.uno("SELECT detalle->>'canal' FROM auditoria WHERE accion = 'CLIENTE_REGISTRADO'") == "PUBLICO"


def test_registro_publico_repetido_responde_igual_y_no_sobrescribe(api):
    primera = api.post("/publico/registro", persona())
    segunda = api.post("/publico/registro", persona(email="otro@example.com", nombres="IMPOSTOR"))
    assert (segunda.status_code, segunda.json()) == (primera.status_code, primera.json())
    fila = api.e.datos.conn.execute(
        "SELECT nombres, email FROM cliente WHERE numero_documento = '1049612345'"
    ).fetchone()
    assert fila == ("LUIS CARLOS", "luis.rojas@example.com")


def test_registro_publico_con_captcha(armar, monkeypatch):
    from estacion.api import captcha as modulo
    from estacion.api.captcha import Captcha

    respuestas = {"bueno": {"success": True}, "malo": {"success": False}}
    enviados: list[dict[str, str]] = []

    def enviar(url, datos, timeout):
        enviados.append(dict(datos))
        return respuestas[datos["response"]]

    monkeypatch.setattr(modulo, "_enviar", enviar)
    api = armar(captcha=Captcha(proveedor="hcaptcha", clave_sitio="sitio-publico", secreto="secreto"))
    captcha = api.get("/publico/configuracion").json()["captcha"]
    assert captcha == {"proveedor": "hcaptcha", "clave_sitio": "sitio-publico"}
    sin_token = api.post("/publico/registro", persona())
    malo = api.post("/publico/registro", persona(captcha="malo"))
    assert sin_token.status_code == malo.status_code == 422
    contar = "SELECT count(*) FROM cliente WHERE numero_documento = '1049612345'"
    assert api.e.datos.uno(contar) == 0
    assert api.post("/publico/registro", persona(captcha="bueno")).status_code == 202
    assert api.e.datos.uno(contar) == 1
    assert [d["response"] for d in enviados] == ["malo", "bueno"]  # sin token no se consulta al proveedor


def test_configuracion_publica_sin_captcha(api):
    assert api.get("/publico/configuracion").json()["captcha"] is None


def test_configuracion_publica_trae_el_nombre_comercial_de_la_estacion_activa(api):
    api.e.datos.conn.execute("UPDATE estacion SET nombre_comercial = 'EDS DE PRUEBA' WHERE id = %s", (api.e.estacion,))
    assert api.get("/publico/configuracion").json()["estacion"] == {"nombre": "EDS DE PRUEBA"}


def test_configuracion_publica_sin_nombre_comercial_usa_la_razon_social(api):
    assert api.get("/publico/configuracion").json()["estacion"] == {"nombre": "ESTACION DE PRUEBA S A S"}


def test_configuracion_publica_sin_estacion_activa(api):
    api.e.datos.conn.execute("UPDATE estacion SET activa = false")
    assert api.get("/publico/configuracion").json()["estacion"] is None


def test_registro_publico_valida_y_limita(armar):
    api = armar(registro_por_minuto=2)
    assert api.post("/publico/registro", persona(autoriza_tratamiento=False)).status_code == 422
    assert api.post("/publico/registro", persona()).status_code == 202
    assert api.post("/publico/registro", empresa()).status_code == 429


# ------------------------------------------------------------------------------------------ facturas
def pedir_factura(api: Api, despacho: int, **cambios: Any):
    datos = {"clave": clave(), "origen_tipo": "despacho", "origen_id": despacho, "cliente_id": api.e.cliente,
             "medio_pago": "EFECTIVO", **cambios}  # fmt: skip
    return api.post("/facturas", datos, api.bombero)


def test_emitir_y_consultar_factura(api):
    despacho = api.e.datos.despacho()
    k = clave()
    r = pedir_factura(api, despacho, clave=k)
    assert r.status_code == 201
    f = r.json()
    assert (f["numero"], f["estado"], f["total"], f["nueva"]) == ("SETP1", "EN_PROCESO", "29757.00", True)
    assert f["cliente"]["nombre_mostrar"] == "ANA MARÍA PÉREZ GÓMEZ" and len(f["cufe"]) == 96

    otra_vez = pedir_factura(api, despacho, clave=k)  # doble clic
    assert otra_vez.status_code == 200 and otra_vez.json()["id"] == f["id"] and otra_vez.json()["nueva"] is False

    procesar_pendientes(api.e.datos.conn, ProveedorSimulado())
    consulta = api.get(f"/facturas/{f['id']}", api.bombero).json()
    assert consulta["estado"] == "FACTURADO" and consulta["nueva"] is None


def test_pdf_solo_de_facturas_validadas_y_queda_en_la_auditoria(api):
    f = pedir_factura(api, api.e.datos.despacho()).json()
    r = api.get(f"/facturas/{f['id']}/pdf", api.bombero)
    assert r.status_code == 409 and "EN_PROCESO" in r.json()["detail"]

    procesar_pendientes(api.e.datos.conn, ProveedorSimulado())
    r = api.get(f"/facturas/{f['id']}/pdf", api.bombero)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-")
    assert r.headers["cache-control"] == "no-store"  # lleva datos personales
    assert r.headers["content-disposition"] == 'inline; filename="factura_SETP1.pdf"'
    auditoria = api.e.datos.conn.execute(
        "SELECT entidad_id, detalle FROM auditoria WHERE accion = 'FACTURA_PDF'"
    ).fetchall()
    assert auditoria == [(str(f["id"]), {})]


def test_pdf_exige_sesion_y_factura_existente(api):
    assert api.get("/facturas/1/pdf").status_code == 401
    assert api.get("/facturas/999999/pdf", api.bombero).status_code == 404


@pytest.mark.parametrize(
    ("respuesta", "texto"),
    [
        ({"simulado": True, "emitida_en": "2026-10-01T10:00:00"}, "SIMULADA (banco de pruebas): no se envió a la DIAN"),
        ({"fecha_validacion": "2026-10-01 10:05:00"}, "2026-10-01 10:05:00"),
        ({}, "SIN DATO DE VALIDACIÓN"),
        (None, "SIN DATO DE VALIDACIÓN"),
    ],
)
def test_texto_de_validacion_nunca_inventa_una_fecha(respuesta, texto):
    from estacion.api.rutas.facturas import texto_validacion

    assert texto_validacion(respuesta) == texto


def test_otra_clave_para_una_venta_en_proceso(api):
    despacho = api.e.datos.despacho()
    pedir_factura(api, despacho)
    r = pedir_factura(api, despacho)
    assert r.status_code == 409 and "no está pendiente" in r.json()["detail"]


@pytest.mark.parametrize(
    ("cambios", "codigo", "texto"),
    [
        ({"medio_pago": "TARJETA"}, 422, "caso todavía no soportado"),
        ({"clave": "corta"}, 422, "clave"),
        ({"cliente_id": 999_999}, 422, "no existe"),
        ({"origen_id": 999_999}, 409, "no existe"),
    ],
)
def test_factura_rechazada_antes_de_numerar(api, cambios, codigo, texto):
    r = pedir_factura(api, api.e.datos.despacho(), **cambios)
    assert r.status_code == codigo and texto in r.text
    assert api.e.siguiente() == 1


def test_comprador_con_nit_todavia_no_se_factura(api):
    cliente = api.post("/clientes", empresa(), api.bombero).json()
    r = pedir_factura(api, api.e.datos.despacho(), cliente_id=cliente["id"])
    assert r.status_code == 422 and "todavía no soportado" in r.json()["detail"]


def test_sin_configuracion_dian_no_se_emite(armar):
    api = armar(entorno_dian={})
    r = pedir_factura(api, api.e.datos.despacho())
    assert r.status_code == 503 and "DIAN_CLAVE_TECNICA" in r.json()["detail"]
    assert ENTORNO_DIAN["DIAN_SOFTWARE_PIN"] not in r.text


def test_factura_inexistente(api):
    assert api.get("/facturas/999999", api.bombero).status_code == 404


# ------------------------------------------------------------------------------------------ reportes
def test_reportes_solo_para_administrador(api):
    for ruta in ("/reportes/diario?desde=2026-09-01&hasta=2026-09-30", "/anomalias"):
        assert api.get(ruta, api.bombero).status_code == 403


def test_reportes_de_ventas(api):
    api.e.datos.despacho()
    api.e.datos.despacho()
    # El turno sintético de Datos.despacho() usa now(): se ubica en el día de los despachos.
    api.e.datos.conn.execute("UPDATE turno SET inicio = '2026-09-30 10:00', fin = '2026-09-30 10:02'")
    rango = "desde=2026-09-01&hasta=2026-09-30"
    (dia,) = api.get(f"/reportes/diario?{rango}", api.admin).json()
    assert (dia["periodo"], dia["despachos"], dia["galones"], dia["valor"]) == ("2026-09-30", 2, "3.722", "59514")
    (mes,) = api.get(f"/reportes/mensual?{rango}", api.admin).json()
    assert (mes["periodo"], mes["valor"]) == ("2026-09-01", "59514")
    (turno,) = api.get(f"/reportes/turno?{rango}", api.admin).json()
    assert turno["despachos"] == 2
    assert api.get("/reportes/diario?desde=2026-09-30&hasta=2026-09-01", api.admin).status_code == 422


def test_reportes_suman_las_manuales_sin_enlace_y_no_duplican_las_enlazadas(api):
    d = api.e.datos
    d.despacho()
    enlazado = d.despacho()
    d.venta_manual(api.e.usuario)  # sin enlace: 3,120 gal por $50.000 el 30-sep
    otra = d.venta_manual(api.e.usuario, volumen="1.861", valor=29757)
    d.conn.execute("UPDATE venta_manual SET despacho_id = %s WHERE id = %s", (enlazado, otra))  # la misma venta
    rango = "desde=2026-09-01&hasta=2026-09-30"
    for ruta, periodo in (("diario", "2026-09-30"), ("mensual", "2026-09-01")):
        filas = {f["origen"]: f for f in api.get(f"/reportes/{ruta}?{rango}", api.admin).json()}
        assert set(filas) == {"SURTIDOR", "MANUAL"}, ruta
        assert (filas["SURTIDOR"]["despachos"], filas["SURTIDOR"]["galones"], filas["SURTIDOR"]["valor"]) == (
            2, "3.722", "59514",
        )  # fmt: skip
        assert (filas["MANUAL"]["despachos"], filas["MANUAL"]["galones"], filas["MANUAL"]["valor"]) == (
            1, "3.120", "50000",
        )  # fmt: skip
        assert filas["MANUAL"]["periodo"] == periodo


def test_posibles_duplicados_manual_sin_enlace_con_su_despacho(api):
    d = api.e.datos
    despacho = d.despacho()  # lado A, DIESEL, 1,861 gal, $29.757, 30-sep 10:00
    igual = dict(lado="A", volumen="1.861", valor=29757, ppu=15990)
    duplicada = d.venta_manual(api.e.usuario, **igual, registrada_en="2026-09-30 10:30")
    d.venta_manual(api.e.usuario, **{**igual, "lado": "B"}, registrada_en="2026-09-30 10:30")  # otro lado
    d.venta_manual(api.e.usuario, **{**igual, "volumen": "2.500", "valor": 39975}, registrada_en="2026-09-30 10:30")
    d.venta_manual(api.e.usuario, **igual, registrada_en="2026-09-30 13:00")  # 3 horas después
    r = api.get("/reportes/posibles-duplicados?desde=2026-09-01&hasta=2026-09-30", api.admin)
    assert r.status_code == 200
    assert [(x["venta_manual_id"], x["despacho_id"]) for x in r.json()] == [(duplicada, despacho)]


def test_posibles_duplicados_ignora_lo_ya_enlazado(api):
    d = api.e.datos
    despacho = d.despacho()
    igual = dict(lado="A", volumen="1.861", valor=29757, ppu=15990, registrada_en="2026-09-30 10:30")
    enlazada = d.venta_manual(api.e.usuario, **igual)
    d.conn.execute("UPDATE venta_manual SET despacho_id = %s WHERE id = %s", (despacho, enlazada))
    d.venta_manual(api.e.usuario, **igual)  # su despacho ya está enlazado a otra: no es un doble conteo con este
    assert api.get("/reportes/posibles-duplicados?desde=2026-09-01&hasta=2026-09-30", api.admin).json() == []


def test_posibles_duplicados_solo_administrador(api):
    assert api.get("/reportes/posibles-duplicados?desde=2026-09-01&hasta=2026-09-30", api.bombero).status_code == 403


def _ventas_para_el_contador(api: Api) -> int:
    """Despacho facturado + despacho sin factura + venta manual sin factura, todos del 30-sep. Devuelve el id de
    la factura."""
    d = api.e.datos
    facturado = d.despacho()
    d.despacho()
    d.venta_manual(api.e.usuario)  # 3,120 gal, $50.000
    f = pedir_factura(api, facturado).json()
    procesar_pendientes(d.conn, ProveedorSimulado())
    return f["id"]


def test_exportacion_contador_excel(api):
    _ventas_para_el_contador(api)
    hoy = date.today().isoformat()
    r = api.get(f"/reportes/contador?desde=2026-09-01&hasta={hoy}&formato=xlsx", api.admin)
    assert r.status_code == 200
    assert r.headers["cache-control"] == "no-store"  # lleva datos personales
    assert "attachment" in r.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Facturas", "Resumen diario", "Notas"]
    encabezado, *facturas = list(libro["Facturas"].iter_rows(values_only=True))
    assert len(facturas) == 1
    fila = dict(zip(encabezado, facturas[0], strict=True))
    documento = api.e.datos.uno("SELECT numero_documento FROM cliente WHERE id = %s", (api.e.cliente,))
    identidad = (fila["Número"], fila["Estado"], fila["Tipo doc."], fila["Documento"])
    assert identidad == ("SETP1", "FACTURADO", "CC", documento)
    assert (fila["Cliente"], Decimal(str(fila["Galones"])), Decimal(str(fila["Total"]))) == (
        "ANA MARÍA PÉREZ GÓMEZ", Decimal("1.861"), Decimal("29757"),
    )  # fmt: skip
    assert (fila["Impuestos"], fila["Medio de pago"], fila["Origen"]) == (0, "EFECTIVO", "SURTIDOR")
    encabezado, *dias = list(libro["Resumen diario"].iter_rows(values_only=True))
    resumen = {str(f[0])[:10]: dict(zip(encabezado, f, strict=True)) for f in dias}
    dia = resumen["2026-09-30"]
    assert (dia["Ventas"], Decimal(str(dia["Valor vendido"]))) == (3, Decimal("109514"))
    assert (Decimal(str(dia["Valor facturado"])), Decimal(str(dia["Valor sin factura"]))) == (
        Decimal("29757"), Decimal("79757"),
    )  # fmt: skip
    auditoria = api.e.datos.conn.execute(
        "SELECT entidad_id, detalle FROM auditoria WHERE accion = 'EXPORTE_CONTADOR'"
    ).fetchall()
    assert auditoria == [(f"2026-09-01/{hoy}", {"desde": "2026-09-01", "hasta": hoy, "formato": "xlsx"})]


def test_exportacion_contador_csv_para_excel_en_espanol(api):
    _ventas_para_el_contador(api)
    r = api.get(f"/reportes/contador?desde=2026-09-01&hasta={date.today().isoformat()}&formato=csv", api.admin)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.content.startswith("\ufeff".encode())  # BOM: Excel lo abre como UTF-8 (tildes bien)
    lineas = r.content.decode("utf-8-sig").splitlines()
    assert lineas[0].split(";")[:3] == ["Fecha", "Número", "CUFE"]
    assert len(lineas) == 2
    assert ";29757,00;" in lineas[1] and ";1,861;" in lineas[1]  # coma decimal


def test_exportacion_contador_solo_administrador_y_formato_valido(api):
    rango = "desde=2026-09-01&hasta=2026-09-30"
    assert api.get(f"/reportes/contador?{rango}&formato=xlsx", api.bombero).status_code == 403
    assert api.get(f"/reportes/contador?{rango}&formato=doc", api.admin).status_code == 422


def test_reporte_pdf_por_api(api):
    api.e.datos.despacho()
    rango = "desde=2026-09-01&hasta=2026-09-30"
    for tipo in ("diario", "mensual", "turno"):
        r = api.get(f"/reportes/pdf?tipo={tipo}&{rango}", api.admin)
        assert r.status_code == 200 and r.headers["content-type"] == "application/pdf", tipo
        assert r.content.startswith(b"%PDF-")
        assert f'filename="ventas_{tipo}_20260901_20260930.pdf"' in r.headers["content-disposition"]
    assert api.get(f"/reportes/pdf?tipo=diario&{rango}", api.bombero).status_code == 403
    assert api.get(f"/reportes/pdf?tipo=anual&{rango}", api.admin).status_code == 422


def test_exportacion_contador_pdf_queda_en_la_auditoria(api):
    _ventas_para_el_contador(api)
    hoy = date.today().isoformat()
    r = api.get(f"/reportes/contador?desde=2026-09-01&hasta={hoy}&formato=pdf", api.admin)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["cache-control"] == "no-store"
    formatos = api.e.datos.conn.execute(
        "SELECT detalle->>'formato' FROM auditoria WHERE accion = 'EXPORTE_CONTADOR'"
    ).fetchall()
    assert formatos == [("pdf",)]


def test_exportar_excel_desde_la_base(api):
    sim = SimuladorSurtidor(semilla=5, reloj=lambda: datetime(2026, 9, 30, 9, 0))
    for _ in range(8):
        sim.despachar()
    ingerir(api.e.datos.bd.dsn, sim)
    surtidor = api.e.datos.uno("SELECT id FROM surtidor WHERE marca = 'SIMULADOR'")
    r = api.get(f"/reportes/exportar?surtidor_id={surtidor}&desde=2026-09-30&hasta=2026-09-30", api.admin)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Diario", "Mensual", "Turnos", "Datos", "Alertas", "Notas"]
    assert libro["Datos"].max_row == 9  # encabezado + 8 despachos
    total = sum(Decimal(str(fila[8])) for fila in libro["Datos"].iter_rows(min_row=2, values_only=True))
    assert total == api.e.datos.uno("SELECT sum(valor) FROM despacho")
    assert api.get("/reportes/exportar?surtidor_id=999&desde=2026-09-30&hasta=2026-09-30", api.admin).status_code == 404


def test_anomalias(api):
    sim = SimuladorSurtidor(semilla=5, reloj=lambda: datetime(2026, 9, 30, 9, 0))
    sim.despachar("A", pistola=1)
    sim.despachar("A", pistola=1, falla="salto")
    ingerir(api.e.datos.bd.dsn, sim)
    altas = api.get("/anomalias?severidad=ALTA", api.admin).json()
    assert [(a["tipo"], a["despacho_id_externo"]) for a in altas] == [("GAP_TOTALIZADOR", 3)]
    assert api.get("/anomalias?estado=resueltas", api.admin).json() == []


# ------------------------------------------------------------------------------------------ flujo completo
def test_flujo_completo_por_api_contra_el_simulador(api):
    """Aceptación de la Fase 4: venta del simulador -> cliente nuevo -> emitir -> worker -> FACTURADO."""
    sim = SimuladorSurtidor(semilla=9)
    sim.despachar("B", pistola=1, galones=Decimal("4.500"), forma_pago="CONTADO")
    ingerir(api.e.datos.bd.dsn, sim)

    cabeceras = api.entrar(api.bombero)
    (venta,) = api.cliente.get("/despachos/pendientes?lado=B", headers=cabeceras).json()
    no_esta = api.cliente.get("/clientes/buscar?tipo=CC&numero=1049612345", headers=cabeceras)
    assert no_esta.status_code == 404
    cliente = api.cliente.post("/clientes", json=persona(), headers=cabeceras).json()
    datos = {"clave": clave(), "origen_tipo": venta["origen_tipo"], "origen_id": venta["origen_id"],
             "cliente_id": cliente["id"], "medio_pago": "EFECTIVO"}  # fmt: skip
    factura = api.cliente.post("/facturas", json=datos, headers=cabeceras).json()
    assert factura["estado"] == "EN_PROCESO" and factura["total"] == f"{venta['valor']}.00"
    assert api.cliente.get("/despachos/pendientes?lado=B", headers=cabeceras).json() == []

    procesar_pendientes(api.e.datos.conn, ProveedorSimulado())  # el worker, en segundo plano
    final = api.cliente.get(f"/facturas/{factura['id']}", headers=cabeceras).json()
    assert final["estado"] == "FACTURADO" and final["cliente"]["nombre_mostrar"] == "LUIS CARLOS ROJAS"


# ------------------------------------------------------------------------------------------ utilidades
def test_el_log_de_accesos_no_guarda_la_consulta():
    registro = logging.LogRecord("uvicorn.access", logging.INFO, "", 0, '%s - "%s %s HTTP/%s" %d', None, None)
    registro.args = ("127.0.0.1:5000", "GET", "/clientes/buscar?tipo=CC&numero=1049612345", "1.1", 200)
    SinConsulta().filter(registro)
    assert "1049612345" not in registro.getMessage() and "/clientes/buscar" in registro.getMessage()


def test_comando_de_usuarios(escenario, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", escenario.datos.bd.dsn)
    monkeypatch.setattr("sys.stdin", io.StringIO("2468\n"))
    assert usuarios.main(["crear", "--usuario", "Cajero.Uno", "--nombre", "Cajero Uno", "--pin-stdin"]) == 0
    fila = escenario.datos.conn.execute("SELECT rol, pin_hash FROM usuario WHERE usuario = 'cajero.uno'").fetchone()
    assert fila[0] == "BOMBERO" and fila[1].startswith("$argon2id$") and "2468" not in fila[1]
    monkeypatch.setattr("sys.stdin", io.StringIO("2468\n"))
    assert usuarios.main(["crear", "--usuario", "cajero.uno", "--nombre", "Otro", "--pin-stdin"]) == 1
    monkeypatch.setattr("sys.stdin", io.StringIO("12\n"))
    assert usuarios.main(["pin", "--usuario", "cajero.uno", "--pin-stdin"]) == 2  # PIN muy corto
    assert usuarios.main(["desactivar", "--usuario", "cajero.uno"]) == 0
    assert usuarios.main(["desbloquear", "--usuario", "nadie"]) == 1
