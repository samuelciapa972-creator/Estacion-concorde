"""Exportación de las ventas de la estación para la contabilidad (Siigo no las ve: las factura este sistema).

FORMATO PROVISIONAL: el contador todavía no ha definido el suyo (ver CLAUDE.md). Mientras tanto:
  * Excel con tres hojas: "Facturas" (una fila por factura del rango, en CUALQUIER estado: nada se esconde),
    "Resumen diario" (lo vendido por día, facturado y sin factura) y "Notas" (alcance y supuestos).
  * CSV con solo las facturas, separador `;`, coma decimal y BOM UTF-8: Excel en español lo abre bien.
Las facturas se filtran por su fecha de EMISIÓN; el resumen, por la fecha de la VENTA. Lleva datos personales de los
clientes: la API lo entrega solo al administrador, sin caché, y deja la descarga en la auditoría.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from ._bd import Conexion
from .clientes import nombre_para_mostrar


@dataclass(frozen=True, slots=True)
class FacturaContable:
    fecha: datetime
    numero: str
    cufe: str | None
    estado: str
    tipo_documento: str
    documento: str
    cliente: str
    producto: str
    galones: Decimal
    base: Decimal
    impuestos: Decimal
    total: Decimal
    forma_pago: str
    medio_pago: str
    origen: str  # SURTIDOR | MANUAL


@dataclass(frozen=True, slots=True)
class DiaContable:
    fecha: date
    ventas: int
    galones: Decimal
    valor_vendido: Decimal
    valor_facturado: Decimal
    valor_sin_factura: Decimal


COLUMNAS_FACTURAS = (
    "Fecha", "Número", "CUFE", "Estado", "Tipo doc.", "Documento", "Cliente", "Producto", "Galones", "Base",
    "Impuestos", "Total", "Forma de pago", "Medio de pago", "Origen",
)  # fmt: skip
COLUMNAS_RESUMEN = ("Fecha", "Ventas", "Galones", "Valor vendido", "Valor facturado", "Valor sin factura")


def facturas(conn: Conexion, desde: date, hasta: date) -> list[FacturaContable]:
    filas = conn.execute(
        """SELECT f.fecha_emision, f.numero, f.cufe, f.estado, c.tipo_documento, c.numero_documento, c.nombres,
                  c.apellidos, p.codigo, coalesce(d.volumen_bruto, v.volumen), f.total, f.forma_pago, f.medio_pago,
                  CASE WHEN f.despacho_id IS NOT NULL THEN 'SURTIDOR' ELSE 'MANUAL' END
           FROM factura f
           JOIN cliente c ON c.id = f.cliente_id
           LEFT JOIN despacho d ON d.id = f.despacho_id
           LEFT JOIN venta_manual v ON v.id = f.venta_manual_id
           JOIN producto p ON p.id = coalesce(d.producto_id, v.producto_id)
           WHERE f.fecha_emision::date BETWEEN %s AND %s
           ORDER BY f.fecha_emision, f.prefijo, f.consecutivo""",
        (desde, hasta),
    ).fetchall()
    # Sin impuestos en este alcance (el generador rechaza líneas con IVA): base = total.
    return [
        FacturaContable(
            fecha=f[0],
            numero=f[1],
            cufe=f[2],
            estado=f[3],
            tipo_documento=f[4],
            documento=f[5],
            cliente=nombre_para_mostrar(f[4], f[6], f[7]),
            producto=f[8],
            galones=f[9],
            base=f[10],
            impuestos=Decimal(0),
            total=f[10],
            forma_pago=f[11],
            medio_pago=f[12],
            origen=f[13],
        )  # fmt: skip
        for f in filas
    ]


def resumen_diario(conn: Conexion, desde: date, hasta: date) -> list[DiaContable]:
    """Lo vendido por día con la misma regla de los reportes (despachos + manuales sin enlace), separando lo que
    tiene factura VALIDADA (FACTURADO). Una venta manual enlazada se factura por el despacho o por ella misma."""
    filas = conn.execute(
        """WITH ventas AS (
               SELECT d.fecha_operativa AS fecha, d.volumen_bruto AS galones, d.valor,
                      EXISTS (SELECT 1 FROM factura f
                              WHERE f.estado = 'FACTURADO'
                                AND (f.despacho_id = d.id OR f.venta_manual_id IN (
                                     SELECT v.id FROM venta_manual v WHERE v.despacho_id = d.id))) AS facturada
               FROM despacho d
               UNION ALL
               SELECT v.registrada_en::date, v.volumen, v.valor,
                      EXISTS (SELECT 1 FROM factura f WHERE f.estado = 'FACTURADO' AND f.venta_manual_id = v.id)
               FROM venta_manual v WHERE v.despacho_id IS NULL)
           SELECT fecha, count(*), sum(galones), sum(valor),
                  coalesce(sum(valor) FILTER (WHERE facturada), 0), coalesce(sum(valor) FILTER (WHERE NOT facturada), 0)
           FROM ventas WHERE fecha BETWEEN %s AND %s
           GROUP BY fecha ORDER BY fecha""",
        (desde, hasta),
    ).fetchall()
    return [DiaContable(*f) for f in filas]


def _fila_factura(f: FacturaContable) -> tuple[Any, ...]:
    return (f.fecha, f.numero, f.cufe or "", f.estado, f.tipo_documento, f.documento, f.cliente, f.producto,
            f.galones, f.base, f.impuestos, f.total, f.forma_pago, f.medio_pago, f.origen)  # fmt: skip


def excel(conn: Conexion, desde: date, hasta: date) -> bytes:
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Facturas"
    hoja.append(COLUMNAS_FACTURAS)
    for f in facturas(conn, desde, hasta):
        hoja.append(_fila_factura(f))
    resumen = libro.create_sheet("Resumen diario")
    resumen.append(COLUMNAS_RESUMEN)
    for d in resumen_diario(conn, desde, hasta):
        resumen.append((d.fecha, d.ventas, d.galones, d.valor_vendido, d.valor_facturado, d.valor_sin_factura))
    for h in (hoja, resumen):
        for celda in h[1]:
            celda.font = Font(bold=True)
        h.freeze_panes = "A2"
    notas = libro.create_sheet("Notas")
    for texto in (
        "FORMATO PROVISIONAL: lo reemplaza el que defina el contador.",
        f"Rango: {desde.isoformat()} a {hasta.isoformat()}.",
        "Facturas: por fecha de emisión, en todos los estados (FACTURADO es la única validada).",
        "Resumen diario: por fecha de la venta; despachos del surtidor + ventas manuales sin enlace.",
        "Impuestos en 0: el tratamiento del combustible lo define el contador (el sistema no factura con IVA).",
        "Contiene datos personales (Ley 1581 de 2012): no reenviar fuera de la contabilidad.",
    ):
        notas.append((texto,))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _texto_csv(valor: object) -> str:
    if isinstance(valor, Decimal):
        return str(valor).replace(".", ",")  # Excel en español: coma decimal
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    return str(valor)


def csv_facturas(conn: Conexion, desde: date, hasta: date) -> bytes:
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(COLUMNAS_FACTURAS)
    for f in facturas(conn, desde, hasta):
        escritor.writerow([_texto_csv(x) for x in _fila_factura(f)])
    return salida.getvalue().encode("utf-8-sig")
